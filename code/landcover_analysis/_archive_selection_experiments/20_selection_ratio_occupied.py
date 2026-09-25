#!/usr/bin/env python3
"""20_selection_ratio_occupied.py  --  occupied-range selection ratios, scored

Recomputes the habitat selection ratio using each species' OCCUPIED-RANGE
availability instead of the UK-wide denominator, then scores the result against
the RSPB Handbook exactly as 16_/18_ did, so the two are directly comparable.

THE CHANGE
----------
Script 15 used a single national denominator for every species:

    p_avail[h]    = avail[h] / sum_h avail[h]           <- same for all species

That penalises range-restricted birds. A grebe on Scottish sea lochs was scored
against Kent arable it will never encounter, and because Saltwater is 1.13% of UK
land, dividing by that national share inflated its SR mechanically. The ceiling
test confirmed the damage: species where the data over-called a marine class got
worse under the correction at 0.667 vs 0.290 otherwise.

Here the denominator is per species, over the cells it actually occupies:

    used[s,h]     = sum over OCCUPIED cells of ( abundance * w_h )
    avail[s,h]    = sum over the SAME cells of ( w_h )
    p_use[s,h]    = used[s,h]  / sum_h used[s,h]
    p_avail[s,h]  = avail[s,h] / sum_h avail[s,h]       <- now species-specific
    SR[s,h]       = p_use / p_avail
    Jacobs' D     = (p_use - p_avail) / (p_use + p_avail - 2*p_use*p_avail)

Numerator and denominator run over the same cells with the same weights, differing
only by the abundance factor -- the "same basis" rule from script 15, applied per
species rather than nationally.

WHAT SR MEANS NOW. Under a national denominator, SR asked "does this bird use
habitat h more than the UK offers?". Under an occupied-range denominator it asks
"WITHIN its range, does it concentrate in h more than that range offers?" That is
a within-range preference, and it is the question the Handbook actually answers.
A species can now score SR ~ 1 everywhere if it uses its range in proportion --
which is informative, not a failure.

OCCUPANCY SENSITIVITY. 19_ emits several occupancy thresholds (quantiles of each
species' own non-zero abundance). All are scored here; the choice is reported
rather than assumed, since a strict >0 rule sweeps in eBird's near-zero tail and
drags the denominator back toward UK-wide.

Inputs:
    outputs/.../selection_corrected/occupied_range/occupied_used_wide.csv
    outputs/.../selection_corrected/occupied_range/occupied_avail_wide.csv
    outputs/.../selection_corrected/aggregate9_presence_corrected.csv   (UK-wide, to compare)
    outputs/habitat_agents/blind_presence_matrix.csv                    (agent reference)
    alignment_adjusted.xlsx                                             (Book = ground truth)

Outputs (new files only, under selection_corrected/occupied_range/):
    occupied_selection_ratio.csv       SR at the chosen threshold
    occupied_jacobs.csv                Jacobs' D at the chosen threshold
    occupied_presence.csv              boolean SR >= 1
    occupied_threshold_scores.csv      score for every occupancy threshold
    occupied_vs_ukwide.csv             per-species before/after comparison
    occupied_range_comparison.png
    OCCUPIED_RANGE_SUMMARY.md
"""

from __future__ import annotations

import os
import re
import sys

import numpy as np
import pandas as pd
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "habitat_agents"))
import _agent_common as ac  # noqa: E402

CODE_ROOT = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
_m = re.search(r"(.*)/\.claude/worktrees/[^/]+$", CODE_ROOT)
DATA_ROOT = _m.group(1) if _m else CODE_ROOT

SEL_DIR = os.path.join(DATA_ROOT, "outputs", "landcover_analysis", "selection_corrected")
OCC_DIR = os.path.join(SEL_DIR, "occupied_range")
USED_CSV = os.path.join(OCC_DIR, "occupied_used_wide.csv")
AVAIL_CSV = os.path.join(OCC_DIR, "occupied_avail_wide.csv")
UKWIDE_PRESENCE = os.path.join(SEL_DIR, "aggregate9_presence_corrected.csv")
AGENT_CSV = os.path.join(DATA_ROOT, "outputs", "habitat_agents",
                         "blind_presence_matrix.csv")
ALIGNMENT_XLSX = os.path.join(DATA_ROOT, "alignment_adjusted.xlsx")

OUT_SR = os.path.join(OCC_DIR, "occupied_selection_ratio.csv")
OUT_JACOBS = os.path.join(OCC_DIR, "occupied_jacobs.csv")
OUT_PRESENCE = os.path.join(OCC_DIR, "occupied_presence.csv")
OUT_THRESH = os.path.join(OCC_DIR, "occupied_threshold_scores.csv")
OUT_COMPARE = os.path.join(OCC_DIR, "occupied_vs_ukwide.csv")
OUT_PNG = os.path.join(OCC_DIR, "occupied_range_comparison.png")
OUT_MD = os.path.join(OCC_DIR, "OCCUPIED_RANGE_SUMMARY.md")

CATEGORIES = ac.CATEGORIES
LABELS = ac.CATEGORY_LABELS
N_CLASSES = len(CATEGORIES)

SR_THRESHOLD = 1.0
MARINE = {"saltwater", "coastal"}


def norm(s: str) -> str:
    return re.sub(r"[^a-z]", "", str(s).lower())


CANON = {norm(LABELS[s]): s for s in CATEGORIES}
CANON.update({norm(s): s for s in CATEGORIES})


def load_book() -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(ALIGNMENT_XLSX, data_only=True)
    ws = wb["alignment"]
    header = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    bc, cc = header.get("Book"), header.get("species_code")
    out = {}
    for r in range(2, ws.max_row + 1):
        code = ws.cell(row=r, column=cc).value
        if not code:
            continue
        cell = ws.cell(row=r, column=bc).value
        if cell is None or str(cell).strip() == "":
            out[str(code).strip()] = None
            continue
        slugs = set()
        for part in str(cell).split(";"):
            k = norm(part)
            if k in CANON:
                slugs.add(CANON[k])
            elif k:
                for cand, slug in CANON.items():
                    if cand and cand in k:
                        slugs.add(slug)
                        break
        out[str(code).strip()] = slugs
    return out


def prf_rows(pred, truth):
    out = []
    for i in range(pred.shape[0]):
        d, b = pred[i], truth[i]
        tp = int((d & b).sum()); fp = int((d & ~b).sum())
        fn = int((~d & b).sum()); tn = int((~d & ~b).sum())
        p = tp / (tp + fp) if (tp + fp) else None
        r = tp / (tp + fn) if (tp + fn) else None
        if p is None and r is None:
            f = None
        else:
            pp, rr = (p or 0.0), (r or 0.0)
            f = 2 * pp * rr / (pp + rr) if (pp + rr) > 0 else 0.0
        out.append((p, r, f, (tp + tn) / N_CLASSES))
    return out


def macro(rows, i):
    v = [r[i] for r in rows if r[i] is not None]
    return float(np.mean(v)) if v else float("nan")


def compute_sr(used: np.ndarray, avail: np.ndarray):
    """Selection ratio + Jacobs' D from paired used/avail matrices."""
    ut = used.sum(axis=1, keepdims=True)
    at = avail.sum(axis=1, keepdims=True)
    with np.errstate(divide="ignore", invalid="ignore"):
        p_use = np.where(ut > 0, used / np.where(ut > 0, ut, 1), np.nan)
        p_av = np.where(at > 0, avail / np.where(at > 0, at, 1), np.nan)
        sr = np.where((p_av > 0), p_use / np.where(p_av > 0, p_av, np.nan), np.nan)
        den = p_use + p_av - 2 * p_use * p_av
        jac = np.where(den > 0, (p_use - p_av) / np.where(den > 0, den, 1), np.nan)
    return sr, jac, p_use, p_av


def main() -> None:
    print("=" * 78)
    print("20_selection_ratio_occupied.py -- occupied-range denominator, scored")
    print("=" * 78)
    print("\nResolved inputs:")
    for lbl, p in [("occupied used", USED_CSV), ("occupied avail", AVAIL_CSV),
                   ("UK-wide presence", UKWIDE_PRESENCE), ("blind agent", AGENT_CSV),
                   ("Book", ALIGNMENT_XLSX)]:
        print(f"  [{'OK' if os.path.exists(p) else 'MISSING':>7}] {lbl:<18} {p}")
    for p in (USED_CSV, AVAIL_CSV, UKWIDE_PRESENCE, ALIGNMENT_XLSX):
        if not os.path.exists(p):
            sys.exit(f"\nERROR: missing {p}\n"
                     "Run 19_occupied_range_availability.R first.")
    print()

    used_all = pd.read_csv(USED_CSV)
    avail_all = pd.read_csv(AVAIL_CSV)
    used_all["species_code"] = used_all["species_code"].astype(str).str.strip()
    avail_all["species_code"] = avail_all["species_code"].astype(str).str.strip()

    book = load_book()
    quantiles = sorted(used_all["occ_quantile"].unique())
    print(f"Occupancy thresholds available: {quantiles}")

    # --- score every occupancy threshold -----------------------------------
    rows = []
    per_q = {}
    for q in quantiles:
        u = used_all[used_all["occ_quantile"] == q].sort_values("species_code")
        a = avail_all[avail_all["occ_quantile"] == q].sort_values("species_code")
        codes = u["species_code"].tolist()
        assert codes == a["species_code"].tolist(), "used/avail species mismatch"

        U = u[CATEGORIES].to_numpy(float)
        A = a[CATEGORIES].to_numpy(float)
        sr, jac, p_use, p_av = compute_sr(U, A)
        pred = np.where(np.isnan(sr), False, sr >= SR_THRESHOLD)

        scored = [i for i, c in enumerate(codes) if book.get(c) is not None]
        TRUTH = np.array([[cat in book[codes[i]] for cat in CATEGORIES]
                          for i in scored], dtype=bool)
        rr = prf_rows(pred[scored], TRUTH)
        rows.append({
            "occ_quantile": q,
            "n_species": len(codes), "n_scored": len(scored),
            "mean_occupied_cells": float(u["n_occupied_cells"].mean()),
            "flags_per_species": float(pred.sum(axis=1).mean()),
            "precision": macro(rr, 0), "recall": macro(rr, 1),
            "f1": macro(rr, 2), "accuracy": macro(rr, 3),
        })
        per_q[q] = dict(codes=codes, sr=sr, jac=jac, pred=pred, u=u,
                        scored=scored, TRUTH=TRUTH)

    thresh_df = pd.DataFrame(rows)
    os.makedirs(OCC_DIR, exist_ok=True)
    thresh_df.to_csv(OUT_THRESH, index=False)

    print("\nOccupancy-threshold sensitivity (scored vs the Handbook):")
    print(f"  {'q':>5} {'occ cells':>10} {'flags/sp':>9} {'Prec':>6} "
          f"{'Rec':>6} {'F1':>6} {'Acc':>6}")
    for _, r in thresh_df.iterrows():
        print(f"  {r['occ_quantile']:>5.2f} {r['mean_occupied_cells']:>10.0f} "
              f"{r['flags_per_species']:>9.2f} {r['precision']:>6.3f} "
              f"{r['recall']:>6.3f} {r['f1']:>6.3f} {r['accuracy']:>6.3f}")

    best_q = float(thresh_df.loc[thresh_df["f1"].idxmax(), "occ_quantile"])
    print(f"\n  Best occupancy quantile by F1: {best_q:.2f}")

    # --- headline comparison at the best threshold -------------------------
    P = per_q[best_q]
    codes, sr, jac, pred = P["codes"], P["sr"], P["jac"], P["pred"]
    scored, TRUTH = P["scored"], P["TRUTH"]

    meta = used_all[used_all["occ_quantile"] == best_q].sort_values("species_code")
    ids = meta[["species_code", "common_name", "red_list_name"]].reset_index(drop=True)

    pd.concat([ids, pd.DataFrame(sr, columns=CATEGORIES)], axis=1).to_csv(OUT_SR, index=False)
    pd.concat([ids, pd.DataFrame(jac, columns=CATEGORIES)], axis=1).to_csv(OUT_JACOBS, index=False)
    pd.concat([ids, pd.DataFrame(pred, columns=CATEGORIES)], axis=1).to_csv(OUT_PRESENCE, index=False)

    # UK-wide baseline on the same scored species
    uk = pd.read_csv(UKWIDE_PRESENCE)
    uk["species_code"] = uk["species_code"].astype(str).str.strip()
    uk_idx = uk.set_index("species_code")
    UK = np.array([[bool(uk_idx.loc[codes[i], c]) if codes[i] in uk_idx.index else False
                    for c in CATEGORIES] for i in scored], dtype=bool)

    rows_occ = prf_rows(pred[scored], TRUTH)
    rows_uk = prf_rows(UK, TRUTH)

    agent_rows = None
    if os.path.exists(AGENT_CSV):
        ag = pd.read_csv(AGENT_CSV)
        ag["species_code"] = ag["species_code"].astype(str).str.strip()
        ag_idx = ag.set_index("species_code")
        AG = np.array([[bool(ag_idx.loc[codes[i], c]) if codes[i] in ag_idx.index else False
                        for c in CATEGORIES] for i in scored], dtype=bool)
        agent_rows = prf_rows(AG, TRUTH)

    print("\n" + "=" * 78)
    print(f"HEADLINE  (occupancy q={best_q:.2f}, {len(scored)} species with a Book entry)")
    print("=" * 78)
    print(f"\n  {'Denominator':<34} {'Prec':>6} {'Rec':>6} {'F1':>6} {'Acc':>6}")
    print(f"  {'UK-wide (script 15)':<34} {macro(rows_uk,0):>6.3f} "
          f"{macro(rows_uk,1):>6.3f} {macro(rows_uk,2):>6.3f} {macro(rows_uk,3):>6.3f}")
    print(f"  {'OCCUPIED-RANGE (this script)':<34} {macro(rows_occ,0):>6.3f} "
          f"{macro(rows_occ,1):>6.3f} {macro(rows_occ,2):>6.3f} {macro(rows_occ,3):>6.3f}")
    if agent_rows:
        print(f"  {'blind agent (reference ceiling)':<34} {macro(agent_rows,0):>6.3f} "
              f"{macro(agent_rows,1):>6.3f} {macro(agent_rows,2):>6.3f} "
              f"{macro(agent_rows,3):>6.3f}")
    print(f"  {'prevalence baseline (data-free)':<34} {'—':>6} {'—':>6} "
          f"{0.498:>6.3f} {'—':>6}")

    f_occ = macro(rows_occ, 2); f_uk = macro(rows_uk, 2)
    delta = f_occ - f_uk
    print(f"\n  Change vs UK-wide: {delta:+.3f} macro F1")

    # Paired bootstrap on the improvement.
    rng = np.random.default_rng(0)
    o = np.array([r[2] if r[2] is not None else 0.0 for r in rows_occ])
    w = np.array([r[2] if r[2] is not None else 0.0 for r in rows_uk])
    ns = len(o)
    diffs = np.array([o[idx].mean() - w[idx].mean()
                      for idx in (rng.integers(0, ns, ns) for _ in range(10000))])
    lo, hi = np.percentile(diffs, [2.5, 97.5])
    p_gt = float((diffs > 0).mean())
    print(f"  paired bootstrap 95% CI [{lo:+.3f}, {hi:+.3f}], "
          f"P(occupied > UK-wide) = {p_gt:.3f}")

    # --- did it fix the marine over-calling? -------------------------------
    marine_idx = [CATEGORIES.index(c) for c in MARINE]
    uk_over = UK[:, marine_idx].any(1) & ~TRUTH[:, marine_idx].any(1)
    occ_over = pred[scored][:, marine_idx].any(1) & ~TRUTH[:, marine_idx].any(1)
    print(f"\nMarine over-calling (data claims marine, Book does not):")
    print(f"  UK-wide        : {int(uk_over.sum())} / {len(scored)} species")
    print(f"  occupied-range : {int(occ_over.sum())} / {len(scored)} species")

    # per-class precision/recall shift
    print("\nPer-class precision / recall, UK-wide -> occupied-range:")
    print(f"  {'Class':<24} {'P_uk':>6} {'P_occ':>6}   {'R_uk':>6} {'R_occ':>6}")
    per_class = []
    for j, c in enumerate(CATEGORIES):
        def pr(M):
            tp = int((M[:, j] & TRUTH[:, j]).sum())
            fp = int((M[:, j] & ~TRUTH[:, j]).sum())
            fn = int((~M[:, j] & TRUTH[:, j]).sum())
            return (tp / (tp + fp) if tp + fp else float("nan"),
                    tp / (tp + fn) if tp + fn else float("nan"))
        puk, ruk = pr(UK); pocc, rocc = pr(pred[scored])
        per_class.append({"class": LABELS[c], "p_uk": puk, "p_occ": pocc,
                          "r_uk": ruk, "r_occ": rocc})
        print(f"  {LABELS[c]:<24} {puk:>6.3f} {pocc:>6.3f}   {ruk:>6.3f} {rocc:>6.3f}")

    # --- per-species comparison table --------------------------------------
    comp = []
    for k, i in enumerate(scored):
        comp.append({
            "species_code": codes[i],
            "species": ids.loc[i, "common_name"],
            "book": "; ".join(LABELS[c] for j, c in enumerate(CATEGORIES) if TRUTH[k, j]),
            "ukwide": "; ".join(LABELS[c] for j, c in enumerate(CATEGORIES) if UK[k, j]),
            "occupied": "; ".join(LABELS[c] for j, c in enumerate(CATEGORIES) if pred[i, j]),
            "f1_ukwide": round(w[k], 4), "f1_occupied": round(o[k], 4),
            "delta": round(o[k] - w[k], 4),
        })
    comp_df = pd.DataFrame(comp).sort_values("delta", ascending=False)
    comp_df.to_csv(OUT_COMPARE, index=False)

    n_better = int((comp_df["delta"] > 0).sum())
    n_worse = int((comp_df["delta"] < 0).sum())
    print(f"\nPer species: {n_better} improved, {n_worse} worsened, "
          f"{len(comp_df) - n_better - n_worse} unchanged")
    print("\nBiggest gains:")
    for _, r in comp_df.head(5).iterrows():
        print(f"  {str(r['species'])[:26]:<26} {r['f1_ukwide']:.3f} -> "
              f"{r['f1_occupied']:.3f}  ({r['delta']:+.3f})")
    print("Biggest losses:")
    for _, r in comp_df.tail(5).iterrows():
        print(f"  {str(r['species'])[:26]:<26} {r['f1_ukwide']:.3f} -> "
              f"{r['f1_occupied']:.3f}  ({r['delta']:+.3f})")

    plot_comparison(thresh_df, rows_uk, rows_occ, agent_rows, w, o,
                    per_class, best_q, OUT_PNG)
    write_summary(thresh_df, best_q, rows_uk, rows_occ, agent_rows, delta,
                  lo, hi, p_gt, per_class, comp_df, n_better, n_worse,
                  int(uk_over.sum()), int(occ_over.sum()), len(scored))

    print(f"\nWrote:\n  {OUT_SR}\n  {OUT_JACOBS}\n  {OUT_PRESENCE}\n"
          f"  {OUT_THRESH}\n  {OUT_COMPARE}\n  {OUT_PNG}\n  {OUT_MD}")
    print("\n" + "=" * 78)
    print("RESULT")
    print("=" * 78)
    print(f"  UK-wide denominator      : macro F1 {f_uk:.3f}")
    print(f"  Occupied-range denominator: macro F1 {f_occ:.3f}  ({delta:+.3f})")
    if agent_rows:
        print(f"  blind agent (ceiling)     : macro F1 {macro(agent_rows,2):.3f}")
    print("=" * 78)


def plot_comparison(thresh_df, rows_uk, rows_occ, agent_rows, w, o,
                    per_class, best_q, out_path):
    fig, axes = plt.subplots(1, 3, figsize=(16.5, 5.2))

    ax = axes[0]
    labels = ["UK-wide", "Occupied\nrange"]
    vals = [macro(rows_uk, 2), macro(rows_occ, 2)]
    colors = ["#b0b0b0", "#1f7a4d"]
    if agent_rows:
        labels.append("Blind\nagent"); vals.append(macro(agent_rows, 2))
        colors.append("#4A3AA7")
    bars = ax.bar(labels, vals, color=colors, alpha=.9)
    ax.axhline(0.498, color="black", ls="--", lw=1.2)
    ax.text(len(labels) - .5, .505, "prevalence baseline", fontsize=8, ha="right")
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width() / 2, v + .012, f"{v:.3f}",
                ha="center", fontsize=9, fontweight="bold")
    ax.set_ylabel("Macro F1 vs Handbook"); ax.set_ylim(0, 1)
    ax.set_title("Does the occupied-range denominator help?", fontsize=11)
    ax.grid(axis="y", alpha=.25)

    ax = axes[1]
    ax.scatter(w, o, s=42, color="#1f7a4d", alpha=.75, edgecolor="white", lw=.6)
    ax.plot([0, 1], [0, 1], color="black", ls="--", lw=1.2)
    ax.set_xlabel("UK-wide F1 (per species)")
    ax.set_ylabel("Occupied-range F1 (per species)")
    ax.set_xlim(-.03, 1.03); ax.set_ylim(-.03, 1.03)
    ax.set_title("Per species (above line = occupied-range better)", fontsize=11)
    ax.grid(alpha=.25)

    ax = axes[2]
    ax.plot(thresh_df["occ_quantile"], thresh_df["f1"], "o-",
            color="#1f7a4d", lw=2, label="F1")
    ax.plot(thresh_df["occ_quantile"], thresh_df["precision"], "o--",
            color="#d98032", lw=1.2, ms=4, label="Precision")
    ax.plot(thresh_df["occ_quantile"], thresh_df["recall"], "o--",
            color="#7b52c4", lw=1.2, ms=4, label="Recall")
    ax.axvline(best_q, color="#1f7a4d", ls="-.", lw=1.2)
    ax.set_xlabel("Occupancy quantile (of non-zero abundance)")
    ax.set_ylabel("Macro score"); ax.set_ylim(0, 1)
    ax.set_title("Sensitivity to the occupancy threshold", fontsize=11)
    ax.legend(fontsize=8, frameon=False); ax.grid(alpha=.25)

    fig.suptitle("Occupied-range availability vs the UK-wide denominator",
                 fontsize=12.5, y=1.02)
    fig.tight_layout()
    fig.savefig(out_path, dpi=200, bbox_inches="tight")
    plt.close(fig)


def write_summary(thresh_df, best_q, rows_uk, rows_occ, agent_rows, delta,
                  lo, hi, p_gt, per_class, comp_df, n_better, n_worse,
                  uk_over, occ_over, n_scored):
    L = []
    L.append("# Occupied-range availability: does a per-species denominator help?\n")
    L.append("Script 15 divided each species' habitat use by a **UK-wide** availability "
             "denominator. That penalises range-restricted birds — a grebe on Scottish "
             "sea lochs was scored against Kent arable it will never encounter — and "
             "because Saltwater is 1.13% of UK land, dividing by that national share "
             "inflated its selection ratio mechanically.\n")
    L.append("The ceiling test made the damage concrete: species where the data "
             "over-called a marine class got worse under the correction at a rate of "
             "0.667 vs 0.290 otherwise, while a blind LLM agent reached F1 0.699 "
             "against the Handbook where the data reached 0.501. The headroom was real "
             "and the denominator was the prime suspect.\n")

    L.append("## The change\n")
    L.append("```")
    L.append("used[s,h]    = sum over OCCUPIED cells of ( abundance * w_h )")
    L.append("avail[s,h]   = sum over the SAME cells of ( w_h )        <- no abundance term")
    L.append("p_avail[s,h] = avail[s,h] / sum_h avail[s,h]             <- now per species")
    L.append("SR[s,h]      = p_use / p_avail")
    L.append("```\n")
    L.append("Numerator and denominator run over the same cells with the same weights, "
             "differing only by the abundance factor — the same-basis rule from script "
             "15, applied per species rather than nationally.\n")
    L.append("**What SR means now.** Under a national denominator SR asked *does this "
             "bird use habitat h more than the UK offers?* Under an occupied-range "
             "denominator it asks *within its range, does it concentrate in h more than "
             "that range offers?* — a within-range preference, which is closer to what "
             "the Handbook describes.\n")

    L.append("## Result\n")
    L.append("| Denominator | Precision | Recall | F1 | Accuracy |")
    L.append("|---|---:|---:|---:|---:|")
    L.append(f"| UK-wide (script 15) | {macro(rows_uk,0):.3f} | {macro(rows_uk,1):.3f} | "
             f"{macro(rows_uk,2):.3f} | {macro(rows_uk,3):.3f} |")
    L.append(f"| **Occupied-range** | {macro(rows_occ,0):.3f} | {macro(rows_occ,1):.3f} | "
             f"**{macro(rows_occ,2):.3f}** | {macro(rows_occ,3):.3f} |")
    if agent_rows:
        L.append(f"| Blind agent (reference) | {macro(agent_rows,0):.3f} | "
                 f"{macro(agent_rows,1):.3f} | {macro(agent_rows,2):.3f} | "
                 f"{macro(agent_rows,3):.3f} |")
    L.append(f"| Prevalence baseline (data-free) | — | — | 0.498 | — |")
    L.append("")
    L.append(f"Change vs UK-wide: **{delta:+.3f}** macro F1 (paired bootstrap 95% CI "
             f"[{lo:+.3f}, {hi:+.3f}], P(occupied > UK-wide) = {p_gt:.3f}), over "
             f"{n_scored} species with a Book entry.\n")
    L.append(f"Per species: **{n_better} improved, {n_worse} worsened**.\n")

    L.append("### Marine over-calling\n")
    L.append("Species where the data claims a marine habitat the Handbook does not:\n")
    L.append(f"- UK-wide: **{uk_over}** / {n_scored}")
    L.append(f"- Occupied-range: **{occ_over}** / {n_scored}\n")
    L.append("This was the specific failure the occupied-range denominator was built "
             "to fix, so it is the most direct test of the mechanism.\n")

    L.append("## Occupancy-threshold sensitivity\n")
    L.append("A cell counts as occupied when its annual abundance exceeds a quantile of "
             "the species' own non-zero abundance. A strict `> 0` rule (q = 0) sweeps in "
             "eBird's near-zero tail and drags the denominator back toward UK-wide.\n")
    L.append("| q | Mean occupied cells | Flags/species | Precision | Recall | F1 | Accuracy |")
    L.append("|---:|---:|---:|---:|---:|---:|---:|")
    for _, r in thresh_df.iterrows():
        tag = " **(best)**" if abs(r["occ_quantile"] - best_q) < 1e-9 else ""
        L.append(f"| {r['occ_quantile']:.2f}{tag} | {r['mean_occupied_cells']:.0f} | "
                 f"{r['flags_per_species']:.2f} | {r['precision']:.3f} | "
                 f"{r['recall']:.3f} | {r['f1']:.3f} | {r['accuracy']:.3f} |")
    L.append("")

    L.append("## Per class\n")
    L.append("| Class | Precision (UK-wide → occupied) | Recall (UK-wide → occupied) |")
    L.append("|---|---:|---:|")
    for pc in per_class:
        L.append(f"| {pc['class']} | {pc['p_uk']:.3f} → {pc['p_occ']:.3f} | "
                 f"{pc['r_uk']:.3f} → {pc['r_occ']:.3f} |")
    L.append("")

    L.append("## Biggest per-species changes\n")
    L.append("| Species | Book | UK-wide | Occupied-range | ΔF1 |")
    L.append("|---|---|---|---|---:|")
    for _, r in comp_df.head(8).iterrows():
        L.append(f"| {r['species']} | {r['book'] or '—'} | {r['ukwide'] or '—'} | "
                 f"{r['occupied'] or '—'} | {r['delta']:+.3f} |")
    for _, r in comp_df.tail(4).iterrows():
        L.append(f"| {r['species']} | {r['book'] or '—'} | {r['ukwide'] or '—'} | "
                 f"{r['occupied'] or '—'} | {r['delta']:+.3f} |")
    L.append("")

    L.append("## Why it did not work: the circularity\n")
    L.append("The occupied range is defined by where the bird is, so a specialist's "
             "preferred habitat becomes a large share of its own denominator and "
             "normalises itself away. The Puffin is the clearest case:\n")
    L.append("| Class | Share of Puffin's range | SR (UK-wide) | SR (occupied) |")
    L.append("|---|---:|---:|---:|")
    L.append("| Coastal | 21.3% | 29.26 | 2.21 |")
    L.append("| Saltwater | 6.2% | 2.54 | **0.46** |")
    L.append("| Improved grassland | 34.2% | 0.80 | 0.79 |")
    L.append("")
    L.append("Coastal is 1.61% of the UK but 21.3% of the Puffin's occupied range — "
             "because the Puffin is there. Dividing by that inflated denominator "
             "collapses the signal, and Saltwater drops below the SR ≥ 1 threshold "
             "entirely. The correction removes exactly the association it should keep.\n")
    L.append("This is why F1 *falls monotonically* as the range tightens "
             "(0.493 → 0.487 → 0.480 → 0.422 across q = 0 → 0.5): more restriction "
             "means more circularity, not less bias. The denominator did change "
             "substantially — mean total-variation distance from the UK-wide profile is "
             "0.256 at q = 0, rising to 0.343 at q = 0.5 — so this is a real negative "
             "result, not a threshold that was never tested.\n")
    L.append("The species that *gained* are the ones with a genuine range/UK mismatch "
             "and no single dominant habitat (Red-necked Phalarope +0.400, White-fronted "
             "Goose +0.356). The species that *lost* are the specialists whose habitat "
             "defines their range (Kittiwake −0.400, Puffin −0.333, Goldeneye −0.318).\n")

    L.append("## Caveats\n")
    L.append("- Occupancy is defined from the same eBird abundance surface used in the "
             "numerator, so range and use are not independent. A cell enters the "
             "denominator *because* the bird was modelled there. This is the "
             "circularity above, and it is fatal to the approach as formulated.\n")
    L.append("- The occupancy threshold is selected on the same 40 species it is scored "
             "on. The sensitivity table is reported so the choice can be judged; treat "
             "the best-q figure as optimistic.\n")
    L.append("- eBird Status is a modelled, weather-standardised climatology, so "
             "'occupied' means 'modelled present', not observed.\n")
    L.append("Generated by `code/landcover_analysis/20_selection_ratio_occupied.py` "
             "from `19_occupied_range_availability.R`.")

    with open(OUT_MD, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))


if __name__ == "__main__":
    main()
