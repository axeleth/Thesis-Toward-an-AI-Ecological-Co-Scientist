#!/usr/bin/env python3
"""17_llm_and_data_vs_book_scores.py  --  DATA vs Book and LLM vs Book, side by side

One table with the full confusion counts and derived scores for BOTH predictors,
scored against the same ground truth on the same species:

    | Species | DATA: TP FP FN TN P R F1 Acc | LLM: TP FP FN TN P R F1 Acc |

PREDICTORS
    DATA  the availability-corrected selection ratio, SR >= 1
          (15_selection_ratio_correction.py)
    LLM   the blind habitat agent -- literature knowledge only, never shown the
          data (outputs/habitat_agents/blind_presence_matrix.csv)

GROUND TRUTH
    The RSPB Handbook Book column of alignment_adjusted.xlsx.

    TP  predictor says yes, Book says yes     FP  predictor says yes, Book says no
    FN  predictor says no,  Book says yes     TN  predictor says no,  Book says no

    Precision = TP / (TP + FP)      Recall   = TP / (TP + FN)
    F1        = 2PR / (P + R)       Accuracy = (TP + TN) / 9

Precision is undefined when a predictor flags nothing, recall when the Book lists
nothing. Those are reported as N/A, never as 0 -- "flagged nothing" and "flagged
only wrong things" are different claims.

OVERALL ROWS. Both averages are given because they answer different questions and
can disagree when species differ in habitat count:
    macro  mean of the per-species scores      -- every species counts equally
    micro  pooled TP/FP/FN/TN across species   -- every habitat call counts equally
Micro precision/recall/F1/accuracy are computed from the pooled counts, not by
averaging ratios.

Species with no Book entry are excluded from scoring and listed separately; with
the current focal set all 41 have an entry.

Inputs:
    outputs/landcover_analysis/selection_corrected/aggregate9_presence_corrected.csv
    outputs/habitat_agents/blind_presence_matrix.csv
    alignment_adjusted.xlsx  (Book column)

Outputs:
    outputs/landcover_analysis/selection_corrected/scores_data_vs_llm_per_species.csv
    outputs/landcover_analysis/selection_corrected/scores_data_vs_llm_overall.csv
    outputs/landcover_analysis/selection_corrected/scores_data_vs_llm_per_class.csv
    outputs/landcover_analysis/selection_corrected/SCORES_SUMMARY.md

Usage:
    python code/landcover_analysis/17_llm_and_data_vs_book_scores.py
"""

from __future__ import annotations

import os
import re
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "habitat_agents"))
import _agent_common as ac  # noqa: E402

CODE_ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
_m = re.search(r"(.*)/\.claude/worktrees/[^/]+$", CODE_ROOT)
DATA_ROOT = _m.group(1) if _m else CODE_ROOT

SEL_DIR = os.path.join(DATA_ROOT, "outputs", "landcover_analysis", "selection_corrected")
DATA_CSV = os.path.join(SEL_DIR, "aggregate9_presence_corrected.csv")
LLM_CSV = os.path.join(DATA_ROOT, "outputs", "habitat_agents",
                       "blind_presence_matrix.csv")
ALIGNMENT_XLSX = os.path.join(DATA_ROOT, "alignment_adjusted.xlsx")

OUT_PER_SPECIES = os.path.join(SEL_DIR, "scores_data_vs_llm_per_species.csv")
OUT_OVERALL = os.path.join(SEL_DIR, "scores_data_vs_llm_overall.csv")
OUT_PER_CLASS = os.path.join(SEL_DIR, "scores_data_vs_llm_per_class.csv")
OUT_MD = os.path.join(SEL_DIR, "SCORES_SUMMARY.md")

CATEGORIES = ac.CATEGORIES
LABELS = ac.CATEGORY_LABELS
N_CLASSES = len(CATEGORIES)


def norm(s: str) -> str:
    return re.sub(r"[^a-z]", "", str(s).lower())


CANON = {norm(LABELS[s]): s for s in CATEGORIES}
CANON.update({norm(s): s for s in CATEGORIES})


def load_book() -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(ALIGNMENT_XLSX, data_only=True)
    ws = wb["alignment"]
    header = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    bc, cc = header.get("Book"), header.get("species_code")
    if not bc or not cc:
        sys.exit(f"ERROR: need 'Book' and 'species_code' columns; found {sorted(header)}")
    out = {}
    for r in range(2, ws.max_row + 1):
        code = ws.cell(row=r, column=cc).value
        if not code:
            continue
        cell = ws.cell(row=r, column=bc).value
        if cell is None or str(cell).strip() == "":
            out[str(code).strip()] = None
            continue
        slugs = set()
        for part in str(cell).split(";"):
            k = norm(part)
            if k in CANON:
                slugs.add(CANON[k])
            elif k:
                for cand, slug in CANON.items():
                    if cand and cand in k:
                        slugs.add(slug)
                        break
                else:
                    print(f"  WARNING: unrecognised Book fragment {part.strip()!r}")
        out[str(code).strip()] = slugs
    return out


def confusion(pred: set, truth: set) -> tuple[int, int, int, int]:
    tp = len(pred & truth)
    fp = len(pred - truth)
    fn = len(truth - pred)
    tn = N_CLASSES - tp - fp - fn
    return tp, fp, fn, tn


def scores(tp: int, fp: int, fn: int, tn: int):
    """(precision, recall, f1, accuracy); None where undefined."""
    p = tp / (tp + fp) if (tp + fp) else None
    r = tp / (tp + fn) if (tp + fn) else None
    if p is None or r is None or (p + r) == 0:
        f = 0.0 if (p is not None or r is not None) else None
    else:
        f = 2 * p * r / (p + r)
    return p, r, f, (tp + tn) / N_CLASSES


def fmt(v) -> str:
    return "N/A" if v is None or (isinstance(v, float) and np.isnan(v)) else f"{v:.3f}"


def load_pred(path: str, label: str) -> dict:
    df = pd.read_csv(path)
    df["species_code"] = df["species_code"].astype(str).str.strip()
    missing = [c for c in CATEGORIES if c not in df.columns]
    if missing:
        sys.exit(f"ERROR: {label} matrix missing columns: {', '.join(missing)}")
    return {r["species_code"]: {c for c in CATEGORIES if bool(r[c])}
            for _, r in df.iterrows()}


def main() -> None:
    print("=" * 78)
    print("17_llm_and_data_vs_book_scores.py -- DATA and LLM scored on the Handbook")
    print("=" * 78)
    print("\nResolved inputs:")
    for lbl, p in [("DATA (SR >= 1)", DATA_CSV), ("LLM (blind agent)", LLM_CSV),
                   ("Book (ground truth)", ALIGNMENT_XLSX)]:
        print(f"  [{'OK' if os.path.exists(p) else 'MISSING':>7}] {lbl:<20} {p}")
    for p in (DATA_CSV, LLM_CSV, ALIGNMENT_XLSX):
        if not os.path.exists(p):
            sys.exit(f"\nERROR: missing {p}")
    print()

    data_sets = load_pred(DATA_CSV, "DATA")
    llm_sets = load_pred(LLM_CSV, "LLM")
    book = load_book()

    roster = pd.read_csv(DATA_CSV)
    roster["species_code"] = roster["species_code"].astype(str).str.strip()

    rows, unscored, no_llm = [], [], []
    for _, r in roster.iterrows():
        code = r["species_code"]
        name = r.get("common_name", code)
        b = book.get(code)
        if b is None:
            unscored.append(name)
            continue
        if code not in llm_sets:
            no_llm.append(name)
        d = data_sets.get(code, set())
        l = llm_sets.get(code, set())

        dtp, dfp, dfn, dtn = confusion(d, b)
        ltp, lfp, lfn, ltn = confusion(l, b)
        dp, dr, df1, dacc = scores(dtp, dfp, dfn, dtn)
        lp, lr, lf1, lacc = scores(ltp, lfp, lfn, ltn)

        rows.append({
            "species_code": code, "species": name,
            "book_habitats": "; ".join(LABELS[c] for c in CATEGORIES if c in b),
            "data_habitats": "; ".join(LABELS[c] for c in CATEGORIES if c in d),
            "llm_habitats": "; ".join(LABELS[c] for c in CATEGORIES if c in l),
            "data_TP": dtp, "data_FP": dfp, "data_FN": dfn, "data_TN": dtn,
            "data_precision": dp, "data_recall": dr, "data_f1": df1,
            "data_accuracy": dacc,
            "llm_TP": ltp, "llm_FP": lfp, "llm_FN": lfn, "llm_TN": ltn,
            "llm_precision": lp, "llm_recall": lr, "llm_f1": lf1,
            "llm_accuracy": lacc,
        })

    df = pd.DataFrame(rows)
    n = len(df)
    if unscored:
        print(f"NOTE: {len(unscored)} species have no Book entry and are excluded: "
              f"{', '.join(unscored)}")
    if no_llm:
        print(f"NOTE: {len(no_llm)} species absent from the LLM matrix "
              f"(counted as flagging nothing): {', '.join(no_llm)}")
    print(f"Scored {n} species x {N_CLASSES} classes = {n * N_CLASSES} habitat calls.\n")

    # --- per-species CSV ---------------------------------------------------
    out = df.copy()
    for c in [c for c in out.columns if c.endswith(("precision", "recall", "f1", "accuracy"))]:
        out[c] = out[c].map(lambda v: "N/A" if v is None or pd.isna(v) else f"{v:.4f}")
    out.to_csv(OUT_PER_SPECIES, index=False)

    # --- overall -----------------------------------------------------------
    overall = []
    for tag in ("data", "llm"):
        TP = int(df[f"{tag}_TP"].sum()); FP = int(df[f"{tag}_FP"].sum())
        FN = int(df[f"{tag}_FN"].sum()); TN = int(df[f"{tag}_TN"].sum())
        mip, mir, mif, _ = scores(TP, FP, FN, TN)
        miacc = (TP + TN) / (n * N_CLASSES)          # pooled, not per-species /9

        def mac(col):
            v = [x for x in df[f"{tag}_{col}"] if x is not None and not pd.isna(x)]
            return float(np.mean(v)) if v else float("nan")

        overall.append({
            "predictor": "DATA (SR>=1)" if tag == "data" else "LLM (blind agent)",
            "n_species": n, "n_calls": n * N_CLASSES,
            "TP": TP, "FP": FP, "FN": FN, "TN": TN,
            "macro_precision": mac("precision"), "macro_recall": mac("recall"),
            "macro_f1": mac("f1"), "macro_accuracy": mac("accuracy"),
            "micro_precision": mip, "micro_recall": mir,
            "micro_f1": mif, "micro_accuracy": miacc,
        })
    ov = pd.DataFrame(overall)
    ov.to_csv(OUT_OVERALL, index=False)

    print("=" * 78)
    print("OVERALL")
    print("=" * 78)
    print(f"\n  {'Predictor':<20} {'TP':>4} {'FP':>4} {'FN':>4} {'TN':>4}   "
          f"{'macroP':>7} {'macroR':>7} {'macroF1':>7} {'macroAcc':>8}")
    for _, r in ov.iterrows():
        print(f"  {r['predictor']:<20} {r['TP']:>4} {r['FP']:>4} {r['FN']:>4} "
              f"{r['TN']:>4}   {r['macro_precision']:>7.3f} {r['macro_recall']:>7.3f} "
              f"{r['macro_f1']:>7.3f} {r['macro_accuracy']:>8.3f}")
    print(f"\n  {'Predictor':<20} {'':>4} {'':>4} {'':>4} {'':>4}   "
          f"{'microP':>7} {'microR':>7} {'microF1':>7} {'microAcc':>8}")
    for _, r in ov.iterrows():
        print(f"  {r['predictor']:<20} {'':>4} {'':>4} {'':>4} {'':>4}   "
              f"{r['micro_precision']:>7.3f} {r['micro_recall']:>7.3f} "
              f"{r['micro_f1']:>7.3f} {r['micro_accuracy']:>8.3f}")

    # --- per class ---------------------------------------------------------
    B = np.array([[c in book[r["species_code"]] for c in CATEGORIES]
                  for _, r in df.iterrows()], dtype=bool)
    D = np.array([[c in data_sets.get(r["species_code"], set()) for c in CATEGORIES]
                  for _, r in df.iterrows()], dtype=bool)
    Lm = np.array([[c in llm_sets.get(r["species_code"], set()) for c in CATEGORIES]
                   for _, r in df.iterrows()], dtype=bool)

    pc_rows = []
    for j, c in enumerate(CATEGORIES):
        rec = {"class": LABELS[c], "n_book_true": int(B[:, j].sum())}
        for tag, M in (("data", D), ("llm", Lm)):
            tp = int((M[:, j] & B[:, j]).sum()); fp = int((M[:, j] & ~B[:, j]).sum())
            fn = int((~M[:, j] & B[:, j]).sum()); tn = int((~M[:, j] & ~B[:, j]).sum())
            p, r_, f, _ = scores(tp, fp, fn, tn)
            rec.update({f"{tag}_TP": tp, f"{tag}_FP": fp, f"{tag}_FN": fn,
                        f"{tag}_TN": tn, f"{tag}_precision": p, f"{tag}_recall": r_,
                        f"{tag}_f1": f, f"{tag}_accuracy": (tp + tn) / n})
        pc_rows.append(rec)
    pc = pd.DataFrame(pc_rows)
    pc.to_csv(OUT_PER_CLASS, index=False)

    print("\n" + "=" * 78)
    print("PER CLASS")
    print("=" * 78)
    print(f"\n  {'Class':<24} {'DATA TP/FP/FN/TN':>18} {'P':>6} {'R':>6} "
          f"{'F1':>6} | {'LLM TP/FP/FN/TN':>18} {'P':>6} {'R':>6} {'F1':>6}")
    for _, r in pc.iterrows():
        d_c = f"{r['data_TP']}/{r['data_FP']}/{r['data_FN']}/{r['data_TN']}"
        l_c = f"{r['llm_TP']}/{r['llm_FP']}/{r['llm_FN']}/{r['llm_TN']}"
        print(f"  {r['class']:<24} {d_c:>18} {fmt(r['data_precision']):>6} "
              f"{fmt(r['data_recall']):>6} {fmt(r['data_f1']):>6} | "
              f"{l_c:>18} {fmt(r['llm_precision']):>6} "
              f"{fmt(r['llm_recall']):>6} {fmt(r['llm_f1']):>6}")

    # --- per species to stdout --------------------------------------------
    print("\n" + "=" * 78)
    print("PER SPECIES")
    print("=" * 78)
    print(f"\n  {'Species':<28} {'DATA TP/FP/FN/TN':>17} {'P':>5} {'R':>5} {'F1':>5} "
          f"{'Acc':>5} | {'LLM TP/FP/FN/TN':>16} {'P':>5} {'R':>5} {'F1':>5} {'Acc':>5}")
    for _, r in df.iterrows():
        d_c = f"{r['data_TP']}/{r['data_FP']}/{r['data_FN']}/{r['data_TN']}"
        l_c = f"{r['llm_TP']}/{r['llm_FP']}/{r['llm_FN']}/{r['llm_TN']}"
        print(f"  {str(r['species'])[:28]:<28} {d_c:>17} {fmt(r['data_precision']):>5} "
              f"{fmt(r['data_recall']):>5} {fmt(r['data_f1']):>5} "
              f"{fmt(r['data_accuracy']):>5} | {l_c:>16} {fmt(r['llm_precision']):>5} "
              f"{fmt(r['llm_recall']):>5} {fmt(r['llm_f1']):>5} "
              f"{fmt(r['llm_accuracy']):>5}")

    write_summary(df, ov, pc, n, unscored)

    print(f"\nWrote:\n  {OUT_PER_SPECIES}\n  {OUT_OVERALL}\n  {OUT_PER_CLASS}\n  {OUT_MD}")


def write_summary(df, ov, pc, n, unscored):
    L = []
    L.append("# Data and LLM scored against the RSPB Handbook\n")
    L.append("Full confusion counts and derived scores for both predictors against the "
             "same ground truth, on the same species.\n")
    L.append("| | Book: yes | Book: no |")
    L.append("|---|---|---|")
    L.append("| **Predictor: yes** | TP | FP |")
    L.append("| **Predictor: no** | FN | TN |\n")
    L.append("```")
    L.append("Precision = TP / (TP + FP)     Recall   = TP / (TP + FN)")
    L.append("F1        = 2PR / (P + R)      Accuracy = (TP + TN) / 9")
    L.append("```\n")
    L.append("**DATA** = availability-corrected selection ratio, SR ≥ 1. "
             "**LLM** = blind habitat agent, literature knowledge only, never shown "
             "the data.\n")
    if unscored:
        L.append(f"Excluded for having no Book entry: {', '.join(unscored)}.\n")
    else:
        L.append(f"All {n} focal species have a Book entry; none excluded.\n")

    L.append("## Overall\n")
    L.append("| Predictor | Averaging | TP | FP | FN | TN | Precision | Recall | F1 | Accuracy |")
    L.append("|---|---|---:|---:|---:|---:|---:|---:|---:|---:|")
    for _, r in ov.iterrows():
        L.append(f"| **{r['predictor']}** | macro | {r['TP']} | {r['FP']} | {r['FN']} | "
                 f"{r['TN']} | {r['macro_precision']:.3f} | {r['macro_recall']:.3f} | "
                 f"{r['macro_f1']:.3f} | {r['macro_accuracy']:.3f} |")
        L.append(f"| {r['predictor']} | micro | {r['TP']} | {r['FP']} | {r['FN']} | "
                 f"{r['TN']} | {r['micro_precision']:.3f} | {r['micro_recall']:.3f} | "
                 f"{r['micro_f1']:.3f} | {r['micro_accuracy']:.3f} |")
    L.append("")
    L.append(f"Macro averages the per-species scores ({n} species count equally); micro "
             f"pools the counts ({n * N_CLASSES} habitat calls count equally). They "
             "diverge when species differ in how many habitats they use.\n")

    L.append("## Per class\n")
    L.append("| Class | Book-true | DATA TP/FP/FN/TN | DATA P | DATA R | DATA F1 | "
             "LLM TP/FP/FN/TN | LLM P | LLM R | LLM F1 |")
    L.append("|---|---:|:--:|---:|---:|---:|:--:|---:|---:|---:|")
    for _, r in pc.iterrows():
        L.append(f"| {r['class']} | {r['n_book_true']} | "
                 f"{r['data_TP']}/{r['data_FP']}/{r['data_FN']}/{r['data_TN']} | "
                 f"{fmt(r['data_precision'])} | {fmt(r['data_recall'])} | "
                 f"{fmt(r['data_f1'])} | "
                 f"{r['llm_TP']}/{r['llm_FP']}/{r['llm_FN']}/{r['llm_TN']} | "
                 f"{fmt(r['llm_precision'])} | {fmt(r['llm_recall'])} | "
                 f"{fmt(r['llm_f1'])} |")
    L.append("")

    L.append("## Per species\n")
    L.append("| Species | Book | DATA TP/FP/FN/TN | DATA P | DATA R | DATA F1 | DATA Acc | "
             "LLM TP/FP/FN/TN | LLM P | LLM R | LLM F1 | LLM Acc |")
    L.append("|---|---|:--:|---:|---:|---:|---:|:--:|---:|---:|---:|---:|")
    for _, r in df.iterrows():
        L.append(f"| {r['species']} | {r['book_habitats'] or '—'} | "
                 f"{r['data_TP']}/{r['data_FP']}/{r['data_FN']}/{r['data_TN']} | "
                 f"{fmt(r['data_precision'])} | {fmt(r['data_recall'])} | "
                 f"{fmt(r['data_f1'])} | {fmt(r['data_accuracy'])} | "
                 f"{r['llm_TP']}/{r['llm_FP']}/{r['llm_FN']}/{r['llm_TN']} | "
                 f"{fmt(r['llm_precision'])} | {fmt(r['llm_recall'])} | "
                 f"{fmt(r['llm_f1'])} | {fmt(r['llm_accuracy'])} |")
    L.append("")
    L.append("Generated by `code/landcover_analysis/17_llm_and_data_vs_book_scores.py`.")

    with open(OUT_MD, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))


if __name__ == "__main__":
    main()
